#!/usr/bin/env python3
"""Profile authorized XLSX snapshots locally; export only aggregate evidence and sanitized copies.
Private manifest and HMAC key stay in ignored work/. Never prints cell values.
"""
import argparse, collections, datetime as dt, hashlib, hmac, json, os, re, warnings
from pathlib import Path
import openpyxl
from openpyxl.utils import get_column_letter

ROOT=Path(__file__).resolve().parents[1]
SAFE_CATEGORIES={'p','a','h','l','yes','no','na','n/a','nil','active','inactive','dropped out','dropout','home program','home programme','drop out','school','male','female','ei','fa','pvsw','nursery','pt','ot','speech','vision','adl','ei group','ei one on one','m','f','m-2','f / m','m/f','m\nf','-','not applicable'}
IDENTITY=re.compile(r'name|aadh|adhar|address|phone|contact|email|udid|niramaya|admission|caregiver id|child id',re.I)
ORG=re.compile(r'(?!)')
DATE_HINT=re.compile(r'date|\bdob\b|validity',re.I)

def norm(v): return re.sub(r'\s+',' ',str(v).strip()).casefold()
def label(v): return re.sub(r'[^a-z]','',str(v).lower())
def present(v): return v is not None and (not isinstance(v,str) or bool(v.strip()))
def formula(c): return c.data_type=='f' or type(c.value).__name__=='ArrayFormula'
def digest(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def canonical(v):
 if isinstance(v,(dt.datetime,dt.date)):return v.isoformat()
 if isinstance(v,(str,int,float,bool)) or v is None:return v
 return '[FORMULA]' if type(v).__name__=='ArrayFormula' else '['+type(v).__name__+']'
def field_name(c):
 if isinstance(c.value,(dt.datetime,dt.date)):return 'DATE_COLUMN_'+get_column_letter(c.column)
 if formula(c):return 'DERIVED_DATE_'+get_column_letter(c.column)
 if c.value is None:return 'UNLABELED_'+get_column_letter(c.column)
 return ORG.sub('[NGO]',str(c.value)).strip().replace('\n',' ')
def find_header(ws):
 for row in ws.iter_rows(min_row=1,max_row=min(15,ws.max_row)):
  vals=[label(c.value) for c in row if isinstance(c.value,str) and not formula(c)]
  if any(v in {'childname','childnames','nameofchild','nameofthechild','fullnameofthechild','caregivername','totaldayspresent','slno','sno'} for v in vals):return row[0].row
 return None

OVERRIDES={('S05',8):(1,'support_event_summary'),('S07',1):(1,'roster_or_other'),('S07',3):(1,'visit_activity_log'),('S08',11):(1,'visit_activity_log'),('S08',13):(1,'visit_activity_log'),('S09',12):(1,'visit_activity_log'),('S10',1):(1,'roster_or_other'),('S10',12):(1,'visit_activity_log'),('S10',13):(2,'visit_activity_log')}

def classify(alias,idx,ws,hr):
 if alias=='S01':return {1:'questionnaire_primary',5:'questionnaire_subset',6:'questionnaire_view',7:'questionnaire_view'}.get(idx,'supporting_summary')
 if alias=='S03':return {1:'child_master',2:'fees_view',3:'sibling_roster',4:'child_supplement',5:'historical_visits',6:'child_supplement',10:'child_supplement',12:'severity_view'}.get(idx,'supporting_unclassified')
 if alias=='S05':return 'caregiver_link' if idx==7 else ('support_event_summary' if idx==8 else 'support_attendance')
 if alias=='S12':return 'adl_visit_slots' if idx<=10 else ('adl_annual_summary' if idx==11 else 'adl_goals_summary')
 if hr and any('totaldayspresent'==label(c.value) for c in ws[hr]):return 'daily_attendance'
 if hr and any(label(c.value)=='perchildtotalintervention' for c in ws[hr]):return 'annual_attendance_summary'
 if hr:return 'roster_or_other'
 return 'supporting_unclassified'

def candidate_rows(ws,hr,kind):
 if hr is None:return []
 if kind in {'support_event_summary','visit_activity_log'}:
  return [r for r in range(hr+1,ws.max_row+1) if any(present(ws.cell(r,c).value) and not formula(ws.cell(r,c)) for c in range(1,min(ws.max_column,7)+1))]
 header={c.column:label(c.value) for c in ws[hr]}
 name_cols=[col for col,v in header.items() if v in {'childname','childnames','nameofchild','fullnameofthechild','nameofthechild','caregivername'}]
 anchor=name_cols[0] if name_cols else 1
 return [r for r in range(hr+1,ws.max_row+1) if present(ws.cell(r,anchor).value) and not formula(ws.cell(r,anchor)) and norm(ws.cell(r,anchor).value) not in {'total','grand total','child name','name of child'}]

def main():
 ap=argparse.ArgumentParser();ap.add_argument('--manifest',type=Path,required=True);ap.add_argument('--output',type=Path,default=ROOT/'docs/evidence/data-profile.json');args=ap.parse_args()
 manifest=json.loads(args.manifest.read_text()); private=args.manifest.parent
 global ORG
 terms_path=private/'redaction-terms.json'
 terms=json.loads(terms_path.read_text()) if terms_path.exists() else []
 ORG=re.compile('|'.join(re.escape(t) for t in terms),re.I) if terms else re.compile(r'(?!)')
 metadata_path=private/'google-sheet-metadata.json'
 metadata={x['source']:x['sheets'] for x in json.loads(metadata_path.read_text())} if metadata_path.exists() else {}
 key_path=private/'pseudonym.key'
 if not key_path.exists():key_path.write_bytes(os.urandom(32));key_path.chmod(0o600)
 key=key_path.read_bytes()
 def token(v):return 'MASK_'+hmac.new(key,str(v).encode(),hashlib.sha256).hexdigest()[:20]
 def sanitize(v,header=False,identity=False):
  if v is None:return None
  if isinstance(v,str) and not v.strip():return v
  if isinstance(v,str) and v.startswith('='):return '[FORMULA]'
  if header and isinstance(v,str):return ORG.sub('[NGO]',v)
  if isinstance(v,str):
   if v.startswith('='):return '[FORMULA]'
   if not identity and norm(v) in SAFE_CATEGORIES:return v
   return token(v)
  if isinstance(v,(dt.datetime,dt.date)):
   return (dt.date(2000,1,1)+dt.timedelta(days=int(hmac.new(key,str(v).encode(),hashlib.sha256).hexdigest()[:6],16)%7300)).isoformat()
  if isinstance(v,(float,int)):
   return token(v) if identity or abs(v)>1000000 else v
  return '[FORMULA]' if type(v).__name__=='ArrayFormula' else token(str(v))
 report={'version':2,'method':'Local raw-value profiling; only aggregates/field labels/aliases exported. No record is discarded or migrated.','sources':[],'sheets':[],'issues':[]}
 issues=report['issues']; raw_names={}; identity_columns={}; worksheet_map={}
 def issue(source,sheet,code,refs,count=None):
  if refs:issues.append({'source':source,'sheet':sheet,'code':code,'count':len(refs) if count is None else count,'references':refs})
 for f in manifest:
  alias=f['alias'];p=Path(f['original_path']);ensure_source_integrity(digest(p)==f['sha256'], 'Original checksum changed')
  with warnings.catch_warnings(record=True) as captured:
   warnings.simplefilter('always');wb=openpyxl.load_workbook(p,data_only=False);cached=openpyxl.load_workbook(p,data_only=True)
  if alias in metadata:ensure_source_integrity(len(wb.worksheets)==len(metadata[alias]), 'Export omitted a worksheet')
  source={'source':alias,'sha256':f['sha256'],'bytes':p.stat().st_size,'worksheets':len(wb.worksheets),'parser_warning_count':len(captured),'rows_by_kind':{},'source_record_rows':0,'sanitized_sha256':{}}
  report['sources'].append(source)
  for idx,ws in enumerate(wb.worksheets,1):
   sid=f'T{idx:02}';hr=find_header(ws);kind=classify(alias,idx,ws,hr)
   if (alias,idx) in OVERRIDES:hr,kind=OVERRIDES[(alias,idx)]
   gid=None
   if alias in metadata:
    props=metadata[alias][idx-1]['properties'];ensure_source_integrity(props['index']==idx-1, 'Native sheet order changed')
    native_key=re.sub(r'[^a-z0-9]','',props['title'].casefold());export_key=re.sub(r'[^a-z0-9]','',ws.title.casefold())
    ensure_source_integrity(props['title']==ws.title or (len(props['title'])>31 and native_key[:20]==export_key[:20]), 'Unrecognized exported sheet title change')
    gid=props['sheetId']
   worksheet_map[alias+'_'+sid]={'original_sheet_name':ws.title,'ordinal':idx,'header_row':hr,'google_sheet_id':gid,'native_sheet_name':props['title'] if alias in metadata else ws.title}
   # S07 uses child labels that vary; Total Days Present reliably identifies the header.
   rows=candidate_rows(ws,hr,kind);rowset=set(rows);cw=cached.worksheets[idx-1]
   heads={c.column:field_name(c) for c in ws[hr]} if hr else {}
   si={'source':alias,'sheet':sid,'kind':kind,'google_sheet_id':gid,'state':ws.sheet_state,'max_row':ws.max_row,'max_column':ws.max_column,'header_row':hr,'record_rows':len(rows),'record_row_refs':rows,'nonempty_rows':0,'blank_rows':0,'formula_only_rows':0,'other_nonempty_rows':0,'nonempty_cells':0,'formulas':0,'excel_errors':0,'merged_ranges':len(ws.merged_cells.ranges),'fields':[]}
   for row in ws.iter_rows():
    nonempty=[c for c in row if present(c.value)]
    if not nonempty:si['blank_rows']+=1
    else:
     si['nonempty_rows']+=1
     if all(formula(c) for c in nonempty):si['formula_only_rows']+=1
     elif row[0].row not in rowset:si['other_nonempty_rows']+=1
    si['nonempty_cells']+=len(nonempty);si['formulas']+=sum(formula(c) for c in nonempty)
   error_refs=[];cache_missing=[]
   for row in ws.iter_rows():
    for c in row:
     if c.data_type=='e' or cw[c.coordinate].data_type=='e':error_refs.append(c.coordinate)
     if formula(c) and cw[c.coordinate].value is None:cache_missing.append(c.coordinate)
   si['excel_errors']=len(error_refs);si['formula_cache_missing']=len(cache_missing)
   issue(alias,sid,'excel_error_or_invalid_date_serial',error_refs)
   issue(alias,sid,'formula_cache_missing',cache_missing)
   name_col=next((col for col,h in heads.items() if label(h) in {'childname','childnames','nameofchild','nameofthechild','fullnameofthechild'}),None)
   if name_col:
    raw_names[(alias,sid)]=[norm(ws.cell(r,name_col).value) for r in rows]
   seen=collections.defaultdict(list)
   for r in rows:seen[json.dumps([canonical(c.value) if type(c.value).__name__!='ArrayFormula' else c.value.text for c in ws[r]],sort_keys=True)].append(r)
   groups=[rs for rs in seen.values() if len(rs)>1]
   si['duplicate_record_groups']=len(groups);si['duplicate_record_excess']=sum(len(g)-1 for g in groups)
   issue(alias,sid,'exact_record_duplicate_excess',[str(r) for g in groups for r in g[1:]])
   if name_col:
    ng=collections.defaultdict(list)
    for r in rows:ng[norm(ws.cell(r,name_col).value)].append(r)
    si['distinct_normalized_child_names']=len(ng);si['repeated_name_groups']=sum(len(v)>1 for v in ng.values())
    issue(alias,sid,'repeated_child_name_rows',[str(r) for g in ng.values() if len(g)>1 for r in g])
   for col,h in heads.items():
    if not rows:continue
    cells=[ws.cell(r,col) for r in rows]; vals=[c.value for c in cells if present(c.value)]; nonformula=[c.value for c in cells if present(c.value) and not formula(c)]
    fs={'column':get_column_letter(col),'label':h,'rows':len(rows),'missing':len(rows)-len(vals),'formulas':sum(formula(c) for c in cells),'types':dict(collections.Counter('formula' if formula(c) else c.data_type if present(c.value) else 'missing' for c in cells)), 'distinct_nonformula':len(set(str(canonical(v)) for v in nonformula))}
    whitespace=[c.coordinate for c in cells if isinstance(c.value,str) and c.value!=c.value.strip() and not formula(c)]
    fs['surrounding_whitespace']=len(whitespace)
    issue(alias,sid,'surrounding_whitespace',whitespace)
    if any(t in label(h) for t in ['program','batch','status','gender','centre','severity','bloodgroup','diagnosis','parentalinvolvement']):
     cats=collections.Counter(str(v) for v in nonformula)
     fs['categories']={str(sanitize(v,identity=bool(IDENTITY.search(h)))):cnt for v,cnt in cats.items()}
     fs['normalized_category_count']=len(set(norm(v) for v in nonformula))
    if DATE_HINT.search(h) and not h.startswith(('DATE_COLUMN','DERIVED_DATE')):
     invalid=[];ambiguous=[]
     for c in cells:
      if not present(c.value) or formula(c):continue
      v=c.value
      if isinstance(v,(dt.datetime,dt.date)):continue
      if isinstance(v,str) and re.fullmatch(r'\d{1,2}[/.-]\d{1,2}[/.-]\d{2,4}',v.strip()):
       parts=re.split(r'[/.-]',v.strip());a,b=int(parts[0]),int(parts[1])
       if a<=12 and b<=12:ambiguous.append(c.coordinate)
       try:dt.date(int(parts[2])+(2000 if len(parts[2])==2 else 0),b,a)
       except ValueError:invalid.append(c.coordinate)
      else:invalid.append(c.coordinate)
     fs['unparsed_or_invalid_date_values']=len(invalid);fs['ambiguous_day_month_strings']=len(ambiguous)
     issue(alias,sid,'unparsed_or_invalid_date',invalid);issue(alias,sid,'ambiguous_day_month_date',ambiguous)
    if 'admission' in label(h) or label(h) in {'childid','caregiverid'}:
     identity_columns[(alias,sid,label(h))]=[norm(v) for v in nonformula]
     counters=collections.Counter(norm(v) for v in nonformula)
     fs['candidate_key_missing_or_formula']=len(rows)-len(nonformula);fs['candidate_key_duplicate_excess']=sum(n-1 for n in counters.values())
    si['fields'].append(fs)
   # Date chronology is assessed locally on typed dates only; text and placeholders stay unresolved.
   if kind=='child_master':
    impossible=[];future=[]
    for r in rows:
     birth=ws.cell(r,10).value;joining=ws.cell(r,9).value
     if isinstance(birth,dt.datetime) and isinstance(joining,dt.datetime) and birth>joining:impossible.append('J'+str(r))
     if isinstance(birth,dt.datetime) and birth.date()>dt.date(2026,10,26):future.append('J'+str(r))
    issue(alias,sid,'birth_after_joining',impossible);issue(alias,sid,'future_birth_date',future)
   # Row-backed attendance marks: don't count roster repeats or blank marks as sessions.
   if kind=='daily_attendance' and hr:
    date_cols=[c.column for c in ws[hr] if isinstance(c.value,(dt.datetime,dt.date)) or formula(c) or (isinstance(c.value,str) and re.search(r'\d.*(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)',c.value,re.I))]
    date_cols=[c for c in date_cols if 7<=c<=37]
    marks=[];blank=0;no_date_marks=[];badheader=[];datekeys=[]
    for col in range(7,min(37,ws.max_column)+1):
     hc=cw.cell(hr,col);v=hc.value
     if col in date_cols:
      if isinstance(v,(dt.datetime,dt.date)):datekeys.append(v.date().isoformat() if isinstance(v,dt.datetime) else v.isoformat())
      else:badheader.append(ws.cell(hr,col).coordinate)
     for r in rows:
      c=ws.cell(r,col)
      if present(c.value) and not formula(c):
       marks.append(c.value)
       if col not in date_cols:no_date_marks.append(c.coordinate)
      elif col in date_cols:blank+=1
    si['attendance']={'date_columns':len(date_cols),'date_min':min(datekeys) if datekeys else None,'date_max':max(datekeys) if datekeys else None,'duplicate_date_columns':len(datekeys)-len(set(datekeys)),'nonempty_marks':len(marks),'blank_date_slots':blank,'marks_without_date_column':len(no_date_marks),'codes':dict(collections.Counter(str(sanitize(v)) for v in marks))}
    issue(alias,sid,'unresolved_date_header',badheader);issue(alias,sid,'mark_without_date_header',no_date_marks)
    if si['attendance']['duplicate_date_columns']:issue(alias,sid,'duplicate_date_headers',[str(hr)],si['attendance']['duplicate_date_columns'])
   if kind.startswith('questionnaire') and rows:
    qstart=5 if idx in [1,5] else 4;qend=31 if idx in [1,5] else 30
    values=[ws.cell(r,c) for r in rows for c in range(qstart,qend+1)]
    invalid=[c.coordinate for c in values if present(c.value) and not formula(c) and (not isinstance(c.value,(int,float)) or not 1<=c.value<=5)]
    si['questionnaire_scores']={'candidate_items':qend-qstart+1,'slots':len(values),'missing':sum(not present(c.value) for c in values),'outside_observed_1_to_5_numeric_scale':len(invalid),'numeric_distribution':dict(collections.Counter(str(c.value) for c in values if isinstance(c.value,(int,float))))}
    issue(alias,sid,'questionnaire_candidate_score_outside_1_to_5',invalid)
   # Export rows to a fresh JSONL representation, stripping original filenames, sheet names, comments, links, properties, pivots and formulas.
   dest=private/'sanitized'/(alias+'_'+sid+'.jsonl');dest.parent.mkdir(exist_ok=True)
   with dest.open('w') as out:
    for row in ws.iter_rows():
     cells={}
     for c in row:
      if c.value is None:continue
      h=heads.get(c.column,'');identity=bool(IDENTITY.search(h)) or (alias=='S03' and re.search(r'diagnos|income|fees|caste|religion|severity|reason|comment|remarks',h,re.I))
      cells[get_column_letter(c.column)]=sanitize(c.value,header=(c.row==hr),identity=identity)
     if cells:out.write(json.dumps({'source':alias,'sheet':sid,'row':row[0].row,'cells':cells},ensure_ascii=False)+'\n')
   dest.chmod(0o600);source['sanitized_sha256'][sid]=digest(dest)
   report['sheets'].append(si);source['source_record_rows']+=len(rows);source['rows_by_kind'][kind]=source['rows_by_kind'].get(kind,0)+len(rows)
  ensure_source_integrity(digest(p)==f['sha256'], 'Original modified while profiling')
  print(alias,'processed',len(wb.worksheets),'worksheets')
 # Compare normalized names locally; export counts only. Matching names are candidate links, never asserted person identities.
 master=collections.Counter(raw_names.get(('S03','T01'),[]));links=[]
 for (source,sheet),names in raw_names.items():
  if source=='S03':continue
  links.append({'source':source,'sheet':sheet,'rows':len(names),'unique_match_to_master_rows':sum(master[n]==1 for n in names),'ambiguous_match_rows':sum(master[n]>1 for n in names),'unmatched_rows':sum(master[n]==0 for n in names)})
 report['name_linkage_candidates']=links
 idlinks=[]
 for source_key,target_key in [(('S01','T01','childid'),('S05','T07','childid')),(('S01','T01','caregiverid'),('S05','T07','caregiverid')),(('S05','T07','childid'),('S03','T01','admissionnumber')),(('S01','T01','childid'),('S03','T01','admissionnumber'))]:
  a=identity_columns.get(source_key,[]);b=collections.Counter(identity_columns.get(target_key,[]))
  idlinks.append({'source_field':':'.join(source_key),'target_field':':'.join(target_key),'source_nonblank_rows':len(a),'target_distinct_values':len(b),'unique_match_rows':sum(b[v]==1 for v in a),'ambiguous_match_rows':sum(b[v]>1 for v in a),'unmatched_rows':sum(b[v]==0 for v in a)})
 report['id_linkage_candidates']=idlinks
 (private/'worksheet-map.json').write_text(json.dumps(worksheet_map,indent=2))
 (private/'worksheet-map.json').chmod(0o600)
 report['totals']={'files':len(report['sources']),'worksheets':len(report['sheets']),'record_rows_including_views':sum(s['record_rows'] for s in report['sheets']),'nonempty_cells':sum(s['nonempty_cells'] for s in report['sheets']),'formulas':sum(s['formulas'] for s in report['sheets']),'excel_error_cells':sum(s['excel_errors'] for s in report['sheets']),'kind_record_rows':dict(collections.Counter())}
 kinds=collections.Counter()
 for s in report['sheets']:kinds[s['kind']]+=s['record_rows']
 report['totals']['kind_record_rows']=dict(kinds)
 args.output.parent.mkdir(parents=True,exist_ok=True);args.output.write_text(json.dumps(report,indent=2,ensure_ascii=False)+'\n')
 print(json.dumps(report['totals'],indent=2))


def ensure_source_integrity(condition,message):
 if not condition:raise ValueError(message)

if __name__=='__main__':main()
