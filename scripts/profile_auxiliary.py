#!/usr/bin/env python3
"""Additional local checks for ID pairs, contact formats, ADL checkboxes and source formulas."""
import argparse,collections,datetime as dt,json,re,warnings
from pathlib import Path
import openpyxl
from openpyxl.formula.tokenizer import Tokenizer
from profile_sources import norm,present,formula
ROOT=Path(__file__).resolve().parents[1]

def main():
 ap=argparse.ArgumentParser();ap.add_argument('--manifest',type=Path,required=True);args=ap.parse_args()
 manifest=json.loads(args.manifest.read_text());profile=json.loads((ROOT/'docs/evidence/data-profile.json').read_text())
 bykey={(s['source'],s['sheet']):s for s in profile['sheets']};books={};result={'id_pair_checks':{},'adl':[],'contacts':[],'source_formula_signatures':[],'attendance_calendar':[]}
 warnings.filterwarnings('ignore',category=UserWarning)
 # Data-only mode is not used to execute formulas; all tests inspect literal input values.
 for f in manifest:books[f['alias']]=openpyxl.load_workbook(f['original_path'],data_only=False,read_only=True)
 master=books['S03'].worksheets[0];master_rows=list(master.iter_rows(values_only=True));registry={norm(r[1]):norm(r[2]) for r in master_rows[1:] if len(r)>2 and present(r[2])}
 link=books['S05'].worksheets[6];pairs={norm(r[1]):norm(r[2]) for r in list(link.iter_rows(values_only=True))[1:] if present(r[1])}
 q=books['S01'].worksheets[0];qrows=list(q.iter_rows(values_only=True));bad_pairs=[];name_mismatch=[]
 for r in bykey[('S01','T01')]['record_row_refs']:
  row=qrows[r-1];caregiver=norm(row[51]);child=norm(row[52])
  if pairs.get(caregiver)!=child:bad_pairs.append(str(r))
  if registry.get(child)!=norm(row[0]):name_mismatch.append(str(r))
 result['id_pair_checks']={'questionnaire_rows':len(bykey[('S01','T01')]['record_row_refs']),'caregiver_child_pair_mismatch_rows':len(bad_pairs),'pair_mismatch_row_refs':bad_pairs,'questionnaire_name_disagrees_with_id_linked_master_rows':len(name_mismatch),'name_disagreement_row_refs':name_mismatch,'master_ids_without_caregiver_link':len(set(registry)-set(pairs.values()))}
 for idx,ws in enumerate(books['S12'].worksheets[:10],1):
  sid=f'T{idx:02}';si=bykey[('S12',sid)];rs=list(ws.iter_rows(values_only=True));checked=0;no_date=[];date_unchecked=[];invalid=[]
  for r in si['record_row_refs']:
   vals=rs[r-1]
   for c in [3,5,7,9,11]:
    flag=vals[c-1] if len(vals)>=c else None;date=vals[c] if len(vals)>c else None
    if flag is True:
     checked+=1
     if not present(date):no_date.append(f'{openpyxl.utils.get_column_letter(c+1)}{r}')
    elif present(date):date_unchecked.append(f'{openpyxl.utils.get_column_letter(c+1)}{r}')
    if present(flag) and not isinstance(flag,bool):invalid.append(f'{openpyxl.utils.get_column_letter(c)}{r}')
  result['adl'].append({'source':'S12','sheet':sid,'record_rows':si['record_rows'],'checked_visit_slots':checked,'checked_without_date':len(no_date),'checked_without_date_refs':no_date,'date_without_checked_flag':len(date_unchecked),'date_without_checked_flag_refs':date_unchecked,'nonboolean_nonblank_flags':len(invalid),'nonboolean_flag_refs':invalid})
 # Format checks are recommendations, not a claim of confirmed business validity.
 for si in profile['sheets']:
  if si['source']!='S03' or si['header_row'] is None:continue
  ws=books['S03'].worksheets[int(si['sheet'][1:])-1];rs=list(ws.iter_rows(values_only=True))
  for f in si['fields']:
   h=f['label'].casefold()
   if not any(word in h for word in ['contact','email']):continue
   if 'address' in h:continue
   col=openpyxl.utils.column_index_from_string(f['column']);refs=[];nonblank=0
   for r in si['record_row_refs']:
    v=rs[r-1][col-1] if len(rs[r-1])>=col else None
    if not present(v):continue
    nonblank+=1
    text=str(int(v)) if isinstance(v,(int,float)) else str(v).strip()
    valid=bool(re.fullmatch(r'[^\s@]+@[^\s@]+\.[^\s@]+',text)) if 'email' in h else bool(re.fullmatch(r'(?:\+?91[ -]?)?[6-9]\d{9}',text))
    if not valid:refs.append(f['column']+str(r))
   result['contacts'].append({'source':'S03','sheet':si['sheet'],'column':f['column'],'label':f['label'],'nonblank':nonblank,'candidate_format_review_count':len(refs),'references':refs})
 # Preserve formula lineage and numeric constants without emitting embedded string literals.
 ws=books['S01'].worksheets[0]
 row=list(ws.iter_rows(min_row=2,max_row=2))[0]
 for c in row:
  if not formula(c) or not isinstance(c.value,str):continue
  try:items=Tokenizer(c.value).items
  except Exception:continue
  constants=sorted({t.value for t in items if t.type=='OPERAND' and t.subtype=='NUMBER'})
  functions=sorted({t.value.rstrip('(') for t in items if t.type=='FUNC' and t.subtype=='OPEN'})
  ranges=sorted({t.value for t in items if t.type=='OPERAND' and t.subtype=='RANGE' and re.fullmatch(r'[A-Za-z]+[0-9]+(?::[A-Za-z]+[0-9]+)?',t.value)})
  result['source_formula_signatures'].append({'source':'S01','sheet':'T01','cell':c.coordinate,'functions':functions,'numeric_constants':constants,'same_sheet_ranges':ranges,'status':'source_formula_observed_not_approved_indicator_definition'})
 months={name.lower():i for i,name in enumerate(['January','February','March','April','May','June','July','August','September','October','November','December'],1)}
 for si in profile['sheets']:
  if si['kind']!='daily_attendance':continue
  ws=books[si['source']].worksheets[int(si['sheet'][1:])-1];rs=list(ws.iter_rows(min_row=1,max_row=si['header_row'],values_only=True))
  year=rs[0][1] if len(rs[0])>1 else None;month=rs[1][1] if len(rs)>1 and len(rs[1])>1 else None
  month=months.get(str(month).strip().lower());mismatches=[]
  if isinstance(year,(float,int)) and month:
   for c,v in enumerate(rs[si['header_row']-1],1):
    if 7<=c<=37 and isinstance(v,(dt.datetime,dt.date)) and (v.year!=int(year) or v.month!=month):mismatches.append(openpyxl.utils.get_column_letter(c)+str(si['header_row']))
  result['attendance_calendar'].append({'source':si['source'],'sheet':si['sheet'],'static_period_metadata_available':bool(isinstance(year,(float,int)) and month),'date_header_period_disagreements':len(mismatches),'references':mismatches})
 out=ROOT/'docs/evidence/additional-checks.json';out.write_text(json.dumps(result,indent=2)+'\n')
 print('Additional checks completed; no source values emitted.')

if __name__=='__main__':main()
