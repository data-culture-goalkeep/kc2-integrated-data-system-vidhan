import unittest
from datetime import date, datetime
from copy import deepcopy
import openpyxl
import tempfile
from pathlib import Path
from datetime import timezone, timedelta
from openpyxl.worksheet.formula import ArrayFormula
from historical_recovery import calendar_date, resolve_cell, verified_identity, calendar_grid
from historical_projection import project
from etl_core import stable_id
from etl_core import canonical
from historical_execution import Coordinator, check_approval
from historical_recovery import sha, W
from unittest.mock import patch


def fixture(kind='daily_attendance'):
    return {'source_file':'DEMO', 'source_sheet':'DEMO', 'source_row':2, 'source_column':'G',
        'source_record_id':stable_id('synthetic-source'), 'verified_child_id':stable_id('synthetic-child'),
        'recovered_date':'2024-02-29', 'date_verification':'ORIGINAL_EVIDENCE',
        'dataset':kind, 'original_value':'P', 'observation_kind':'MARK'}


class RecoveryTests(unittest.TestCase):
    def auth_fixture(self):
        now=datetime.now(timezone.utc)
        m={'version':'original-date-recovery-v1','region':'ap-northeast-1','project_reference':'hmlbdozsqxlwzvjimjwh',
           'code_sha256':{},'batches':[{'index':0,'query_sha256':sha(b'SYNTHETIC SQL'),'request_id':'DEMO','recovery_id':'DEMO-R','target_id':'DEMO-T'}]}
        a={'approved':True,'active':True,'manifest_sha256':sha(canonical(m).encode()),'project_reference':m['project_reference'],
          'approval_reference':'SYNTHETIC-TEST-ONLY','issued_at':(now-timedelta(minutes=1)).isoformat(),'expires_at':(now+timedelta(minutes=5)).isoformat()}
        return m,a

    def test_authorization_binding_and_expiry(self):
        m,a=self.auth_fixture();check_approval(a,m)
        with self.assertRaises(ValueError):check_approval(a,{**m,'region':'other'})
        with self.assertRaises(ValueError):check_approval(a,m,datetime.now(timezone.utc)+timedelta(hours=1))
        with self.assertRaises(ValueError):check_approval({**a,'manifest_sha256':'changed'},m)

    def test_checkpoint_unknown_ack_and_ordering(self):
        m,a=self.auth_fixture()
        with tempfile.TemporaryDirectory() as d:
            directory=Path(d);(directory/'date-batches').mkdir();(directory/'date-batches/0000.sql').write_text('SYNTHETIC SQL')
            c=Coordinator(m,a,directory/'journal');c.append('preflight',manifest_sha256=sha(canonical(m).encode()))
            with patch('historical_execution.W',directory):
                c.dispatch(0)
                with self.assertRaises(ValueError):c.dispatch(0)
                with self.assertRaises(ValueError):c.accept(0,{'isError':True})
                with self.assertRaises(ValueError):c.dispatch(0)

    def test_checkpoint_receipt_integrity(self):
        m,a=self.auth_fixture()
        with tempfile.TemporaryDirectory() as d:
            directory=Path(d);(directory/'date-batches').mkdir();(directory/'date-batches/0000.sql').write_text('SYNTHETIC SQL')
            c=Coordinator(m,a,directory/'journal');c.append('preflight',manifest_sha256=sha(canonical(m).encode()))
            with patch('historical_execution.W',directory):c.dispatch(0)
            c.accept(0,{'structuredContent':[{'receipt':{'recovery_id':'DEMO-R','target_id':'DEMO-T','outcome':'completed'}}]})
            self.assertEqual(c.events()[-1]['kind'],'ack')
            (c.directory/'receipt-0000.json').write_text('{}')
            with self.assertRaises(ValueError):c.events()

    def test_leap_year_and_invalid_days(self):
        self.assertEqual(calendar_date('2024-02-29')[0], date(2024,2,29))
        self.assertIsNone(calendar_date('2023-02-29')[0])
        self.assertIsNone(calendar_date(31,2024,4)[0])
        self.assertIsNone(calendar_date(12)[0])

    def test_text_ambiguity_and_named_month(self):
        self.assertEqual(calendar_date('29 -Feb- 2024')[0],date(2024,2,29))
        self.assertIsNone(calendar_date('03/04/2024')[0])
        self.assertEqual(calendar_date('29/2/2024')[0],date(2024,2,29))
        self.assertIsNone(calendar_date('29/2/24')[0])
        self.assertIsNone(calendar_date('!1-Mar-2024')[0])

    def test_formula_dependencies_rollover_cycle(self):
        ws=openpyxl.Workbook().active
        ws['A1']=datetime(2024,2,28);ws['B1']='=A1+1';ws['C1']='=B1+1'
        self.assertEqual(resolve_cell(ws,'C1')[0],date(2024,3,1))
        ws['A2']='=B2';ws['B2']='=A2'
        self.assertIsNone(resolve_cell(ws,'A2')[0])
        ws['A3']='=TODAY()'
        self.assertIsNone(resolve_cell(ws,'A3')[0])

    def test_exact_observed_array_formula(self):
        ws=openpyxl.Workbook().active;ws['B1']=2024;ws['B2']='February'
        ws['G6']=ArrayFormula(ref='G6:AI6',text='=LET(dt,DATE(B1,MONTH(B2&1),1),IF(LEN(B1),SEQUENCE(1,EOMONTH(dt,0)-dt+1,dt),))')
        self.assertEqual(resolve_cell(ws,'G6')[0],date(2024,2,1))
        ws['B1']=None
        self.assertIsNone(resolve_cell(ws,'G6')[0])

    def test_calendar_grid_requires_explicit_month_year_and_consensus(self):
        ws=openpyxl.Workbook().active;ws['B1']=2024;ws['B2']='February'
        ws['G6']='damaged';ws['H6']=datetime(2024,2,2);ws['I6']=datetime(2024,2,3)
        self.assertEqual(calendar_grid(ws,6)['G6'][0],date(2024,2,1))
        self.assertNotIn('AK6',calendar_grid(ws,6))
        ws['I6']=datetime(2024,3,3)
        self.assertEqual(calendar_grid(ws,6),{})

    def test_identity_requires_exact_consistent_ids(self):
        children={('str','C1'):'uuid1'};caregivers={('str','G1'):'uuid1',('str','G2'):'uuid2'}
        self.assertEqual(verified_identity(None,'G1',children,caregivers)[0],'uuid1')
        self.assertIsNone(verified_identity('C1','G2',children,caregivers)[0])
        self.assertIsNone(verified_identity('C1 ','G1',children,caregivers)[0])
        self.assertIsNone(verified_identity(None,None,children,caregivers)[0])

    def test_projection_gates_and_determinism(self):
        r=fixture();d={'approved':True,'approval_reference':'synthetic-test-only','event_kind':'daily_attendance','service_code':'DEMO','mark_mapping':{'P':'P'}}
        self.assertEqual(project(r),[])
        result=project(r,d)
        self.assertEqual(result,project(r,d));self.assertEqual(len(result),2)
        for field,value in [('verified_child_id',None),('recovered_date',None),('observation_kind','FORMULA'),('date_verification','SOURCE_CONFLICT')]:
            changed=deepcopy(r);changed[field]=value
            with self.assertRaises(ValueError):project(changed,d)

    def test_nullable_adl_and_support_preserve_missing_dates(self):
        r=fixture('adl_visit_slots');r['original_value']=True;r['recovered_date']=None;r['date_verification']='UNRESOLVED'
        d={'approved':True,'approval_reference':'synthetic-only','event_kind':'adl_visit_slots','true_means_completed_visit':True}
        self.assertIsNone(project(r,d)[0]['values']['visit_date'])
        r=fixture('support_attendance');r['recovered_date']=None
        d={'approved':True,'approval_reference':'synthetic-only','event_kind':'support_attendance','event_reference':'DEMO','value_mapping':{'P':'literal reviewed participation'}}
        self.assertIsNone(project(r,d)[0]['values']['event_date'])


if __name__=='__main__':unittest.main()
